from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

src='C:/Users/10049/Desktop/周报-刘仁春-精简版.docx'; out='.report_work/relation_weekly.docx'
d=Document(src)
for el in list(d._element.body):
    if el.tag != qn('w:sectPr'): d._element.body.remove(el)
sec=d.sections[0]; sec.top_margin=Cm(1.8); sec.bottom_margin=Cm(1.6); sec.left_margin=Cm(2); sec.right_margin=Cm(2)
def ft(r,size=10.5,bold=False,color='000000'):
    r.font.name='Times New Roman'; r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体'); r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=RGBColor.from_string(color)
def para(text='',center=False,size=10.5,bold=False):
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.space_after=Pt(4); ft(p.add_run(text),size,bold); return p
def heading(text):
    p=para(text,size=13,bold=True); p.paragraph_format.space_before=Pt(8)
def bullet(text):
    p=para('• '+text); p.paragraph_format.left_indent=Cm(.4); p.paragraph_format.first_line_indent=Cm(-.4)
def cell(c,text,bold=False,fill=None):
    c.text=''; p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; ft(p.add_run(text),9,bold,'FFFFFF' if fill else '000000')
    if fill:
        sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),fill); c._tc.get_or_add_tcPr().append(sh)
heading('科研工作周报')
p=para('刘仁春    2026年10月16日至10月23日',center=True,size=10); p.runs[0].font.color.rgb=RGBColor(79,79,79)
heading('本周做了什么实验')
bullet('在 SSv2 Full 数据集上测试 RELATION MATCH 关系匹配模块。')
bullet('测试单独 RELATION MATCH，以及与 TASK MATCH MULTI VELOCITY 结合的版本。')
bullet('增加 RELATION CONTROL 对照，并用配对评测比较关系分支的实际贡献。')
heading('结果对比')
t=d.add_table(rows=1,cols=4); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
for i,h in enumerate(['实验','1 shot %','5 shot %','判断']): cell(t.rows[0].cells[i],h,True,'3F4E5E')
for row in [('当前最佳基线\nTASK MATCH MULTI VELOCITY','67.426','82.338','最佳'),('单独 RELATION MATCH','66.260','81.452','下降'),('组合 RELATION MATCH','66.440','81.478','下降')]:
    cs=t.add_row().cells
    for i,v in enumerate(row): cell(cs[i],v)
heading('结论')
para('本周 RELATION MATCH 没有超过当前最佳基线。单独关系匹配的 1 shot 和 5 shot 为 66.260% 和 81.452%；与 TASK MATCH MULTI VELOCITY 结合后为 66.440% 和 81.478%，仍低于基线 67.426% 和 82.338%。')
para('配对评测显示，关系分支净贡献约为 -0.008 至 +0.034 个百分点，提升很小，暂不建议加入主模型。')
heading('下周计划')
bullet('继续检查关系匹配模块在不同随机种子下的稳定性。')
bullet('分析关系分支对 1 shot 和 5 shot 的差异。')
bullet('继续以 67.426% 和 82.338% 为基线，仅保留能稳定提升的改动。')
d.save(out); print(out)
